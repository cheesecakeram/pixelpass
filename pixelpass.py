from pathlib import Path
from io import BytesIO
import getpass
import os
import secrets
import shlex
import shutil
import subprocess
import sys
import tempfile
import xml.sax.saxutils as xml_escape

try:
    from openpyxl import Workbook, load_workbook
    from cryptography.exceptions import InvalidTag
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.hazmat.primitives.kdf.scrypt import Scrypt
except ImportError as error:
    print(
        "pixelpass requires openpyxl and cryptography. "
        "Install both before running the script."
    )
    raise SystemExit(1) from error


APP_NAME = "pixelpass"
HEADERS = ["Site or custom name", "Username", "Email", "Password"]
DISCLAIMER = (
    "pixelpass stores all data in a spreadsheet. Anyone who have access "
    "to the spreadsheet or the other pixelpass files can read saved data "
    "if unencrypted. Strongly recommend encrypting your data."
    "if lost the master password data may be unaccessible."
)
ENCRYPTED_MAGIC = b"PIXELPASS-ENC-1\n"
MAX_INPUT_LENGTH = 99999


def get_data_dir():
    """Return Pixelpass's platform-specific data directory."""
    if sys.platform.startswith("win"):
        return Path(os.environ.get("APPDATA", Path.home() / "AppData/Roaming")) / APP_NAME
    if sys.platform == "darwin":
        return Path.home() / "Library" / "Application Support" / APP_NAME
    # Linux and BSD-based systems
    return Path.home() / ".local" / "share" / APP_NAME


DATA_DIR = get_data_dir()
DATA_DIR.mkdir(parents=True, exist_ok=True)

SCRIPT_PATH = Path(__file__).resolve() if "__file__" in globals() else None
IS_FROZEN = bool(getattr(sys, "frozen", False))


def relocate_self():
    """Copy the running script or packaged executable into its data directory."""
    if IS_FROZEN:
        source = Path(sys.executable).resolve()
        destination = DATA_DIR / "pixelpass.exe" if sys.platform.startswith("win") else DATA_DIR / "pixelpass"
    elif SCRIPT_PATH is not None:
        source = SCRIPT_PATH
        destination = DATA_DIR / "pixelpass.py"
    else:
        return None

    try:
        if source != destination.resolve():
            shutil.copy2(source, destination)
    except (OSError, RuntimeError):
        # Keep running from the original location if copying is not possible.
        return source

    return destination


INSTALLED_EXECUTABLE = relocate_self()
if INSTALLED_EXECUTABLE is None:
    INSTALLED_EXECUTABLE = Path(sys.executable).resolve()

SETTINGS_PATH = DATA_DIR / "settings.txt"
VAULT_PATH = DATA_DIR / "vault.xlsx"
ENCRYPTED_VAULT_PATH = DATA_DIR / "vault.enc"

DEFAULT_SETTINGS = {
    "clear_terminal": "false",
    "vault_file": "vault.xlsx",
}

SETTINGS = {}
MASTER_PASSWORD = None
ENCRYPTED_MODE = False


def load_settings():
    settings = DEFAULT_SETTINGS.copy()
    if SETTINGS_PATH.exists():
        try:
            for line in SETTINGS_PATH.read_text(encoding="utf-8").splitlines():
                if "=" in line:
                    key, value = line.split("=", 1)
                    key = key.strip()
                    if key in DEFAULT_SETTINGS:
                        settings[key] = value.strip()
        except OSError:
            pass

    # Keep all data paths fixed within Pixelpass's own directory.
    settings["vault_file"] = "vault.xlsx"
    settings["clear_terminal"] = (
        "true" if settings.get("clear_terminal", "").lower() == "true" else "false"
    )
    return settings


def save_settings(settings):
    lines = [
        f"clear_terminal={settings.get('clear_terminal', 'false')}",
        "vault_file=vault.xlsx",
    ]
    SETTINGS_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


SETTINGS = load_settings()


def _derive_key(password, salt):
    return Scrypt(
        salt=salt,
        length=32,
        n=2**15,
        r=8,
        p=1,
    ).derive(password.encode("utf-8"))


def encrypt_bytes(plaintext, password):
    salt = secrets.token_bytes(16)
    nonce = secrets.token_bytes(12)
    key = _derive_key(password, salt)
    ciphertext = AESGCM(key).encrypt(nonce, plaintext, ENCRYPTED_MAGIC)
    return ENCRYPTED_MAGIC + salt + nonce + ciphertext


def decrypt_bytes(encrypted, password):
    header_length = len(ENCRYPTED_MAGIC)
    minimum_length = header_length + 16 + 12 + 16
    if len(encrypted) < minimum_length or not encrypted.startswith(ENCRYPTED_MAGIC):
        raise ValueError("The encrypted vault has an invalid format.")

    offset = header_length
    salt = encrypted[offset:offset + 16]
    offset += 16
    nonce = encrypted[offset:offset + 12]
    offset += 12
    ciphertext = encrypted[offset:]

    key = _derive_key(password, salt)
    try:
        return AESGCM(key).decrypt(nonce, ciphertext, ENCRYPTED_MAGIC)
    except InvalidTag as error:
        raise ValueError("The master password is incorrect or the vault is damaged.") from error


def atomic_write(path, contents):
    """Write a file atomically, using a temporary file in the same directory."""
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=".pixelpass-",
            suffix=".tmp",
            dir=DATA_DIR,
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
            temp_file.write(contents)
            temp_file.flush()
            os.fsync(temp_file.fileno())

        try:
            os.chmod(temp_path, 0o600)
        except OSError:
            pass
        temp_path.replace(path)
    finally:
        if temp_path is not None and temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def create_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vault"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 32
    sheet.column_dimensions["D"].width = 36
    return workbook


def workbook_to_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def open_vault_workbook():
    if ENCRYPTED_MODE:
        if MASTER_PASSWORD is None:
            raise ValueError("The encrypted vault is locked.")
        encrypted = ENCRYPTED_VAULT_PATH.read_bytes()
        workbook_bytes = decrypt_bytes(encrypted, MASTER_PASSWORD)
        return load_workbook(BytesIO(workbook_bytes))

    if not VAULT_PATH.exists():
        return create_workbook()
    return load_workbook(VAULT_PATH)


def save_vault_workbook(workbook):
    if ENCRYPTED_MODE:
        if MASTER_PASSWORD is None:
            raise ValueError("The encrypted vault is locked.")
        encrypted = encrypt_bytes(workbook_to_bytes(workbook), MASTER_PASSWORD)
        atomic_write(ENCRYPTED_VAULT_PATH, encrypted)
    else:
        workbook.save(VAULT_PATH)


def ensure_vault():
    global ENCRYPTED_MODE

    if ENCRYPTED_VAULT_PATH.exists():
        ENCRYPTED_MODE = True
        return

    ENCRYPTED_MODE = False
    if not VAULT_PATH.exists():
        workbook = create_workbook()
        try:
            workbook.save(VAULT_PATH)
        finally:
            workbook.close()


def unlock_encrypted_vault():
    global MASTER_PASSWORD, ENCRYPTED_MODE

    if not ENCRYPTED_VAULT_PATH.exists():
        return True

    ENCRYPTED_MODE = True
    print("This Pixelpass vault is protected by a master password.")
    while True:
        password = getpass.getpass("Master password (or Ctrl+C to exit): ")
        try:
            plaintext = decrypt_bytes(ENCRYPTED_VAULT_PATH.read_bytes(), password)
            # Confirm the decrypted payload is a readable workbook.
            workbook = load_workbook(BytesIO(plaintext))
            workbook.close()
            MASTER_PASSWORD = password
            return True
        except (ValueError, OSError, Exception) as error:
            print(f"Could not unlock the vault: {error}")
            if not ask_yes_no("Try again?", True):
                return False


def _desktop_escape(value):
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("$", "\\$")
        .replace("`", "\\`")
    )


def install_application_menu_entry():
    """Create a local application-menu launcher where practical."""
    try:
        if sys.platform.startswith("linux") or "bsd" in sys.platform:
            applications_dir = Path.home() / ".local" / "share" / "applications"
            applications_dir.mkdir(parents=True, exist_ok=True)
            desktop_path = applications_dir / "pixelpass.desktop"

            if IS_FROZEN:
                command = f'"{_desktop_escape(INSTALLED_EXECUTABLE)}"'
            else:
                command = (
                    f'"{_desktop_escape(sys.executable)}" '
                    f'"{_desktop_escape(INSTALLED_EXECUTABLE)}"'
                )

            desktop_contents = (
                "[Desktop Entry]\n"
                "Version=1.0\n"
                "Type=Application\n"
                "Name=Pixelpass\n"
                "Comment=Local password manager\n"
                f"Exec={command}\n"
                "Terminal=true\n"
                "Categories=Utility;Security;\n"
            )
            desktop_path.write_text(desktop_contents, encoding="utf-8")
            try:
                desktop_path.chmod(0o755)
            except OSError:
                pass

        elif sys.platform == "darwin":
            applications_dir = Path.home() / "Applications"
            app_dir = applications_dir / "Pixelpass.app"
            contents_dir = app_dir / "Contents"
            macos_dir = contents_dir / "MacOS"
            macos_dir.mkdir(parents=True, exist_ok=True)

            if IS_FROZEN:
                launch_command = f"exec {shlex.quote(str(INSTALLED_EXECUTABLE))} \"$@\""
            else:
                launch_command = (
                    f"exec {shlex.quote(sys.executable)} "
                    f"{shlex.quote(str(INSTALLED_EXECUTABLE))} \"$@\""
                )

            launcher = macos_dir / "pixelpass"
            launcher.write_text("#!/bin/sh\n" + launch_command + "\n", encoding="utf-8")
            launcher.chmod(0o755)

            bundle_id = "local.pixelpass.passwordmanager"
            info_plist = (
                "<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n"
                "<!DOCTYPE plist PUBLIC \"-//Apple//DTD PLIST 1.0//EN\" "
                "\"http://www.apple.com/DTDs/PropertyList-1.0.dtd\">\n"
                "<plist version=\"1.0\"><dict>\n"
                "<key>CFBundleName</key><string>Pixelpass</string>\n"
                "<key>CFBundleDisplayName</key><string>Pixelpass</string>\n"
                "<key>CFBundleExecutable</key><string>pixelpass</string>\n"
                f"<key>CFBundleIdentifier</key><string>{xml_escape.escape(bundle_id)}</string>\n"
                "<key>CFBundlePackageType</key><string>APPL</string>\n"
                "<key>LSUIElement</key><true/>\n"
                "</dict></plist>\n"
            )
            (contents_dir / "Info.plist").write_text(info_plist, encoding="utf-8")

        elif sys.platform.startswith("win"):
            # A packaged/frozen build provides the requested .exe launcher.
            # For a normal .py run, the shortcut launches Python with the script.
            start_menu = (
                Path(os.environ.get("APPDATA", Path.home() / "AppData/Roaming"))
                / "Microsoft" / "Windows" / "Start Menu" / "Programs"
            )
            start_menu.mkdir(parents=True, exist_ok=True)
            shortcut = start_menu / "Pixelpass.lnk"

            if IS_FROZEN:
                target = str(INSTALLED_EXECUTABLE)
                arguments = ""
                working_dir = str(DATA_DIR)
            else:
                target = sys.executable
                arguments = subprocess.list2cmdline([str(INSTALLED_EXECUTABLE)])
                working_dir = str(DATA_DIR)

            def ps_quote(value):
                return "'" + value.replace("'", "''") + "'"

            ps_script = (
                "$s=(New-Object -ComObject WScript.Shell).CreateShortcut("
                + ps_quote(str(shortcut))
                + ");$s.TargetPath="
                + ps_quote(target)
                + ";$s.Arguments="
                + ps_quote(arguments)
                + ";$s.WorkingDirectory="
                + ps_quote(working_dir)
                + ";$s.Save()"
            )
            subprocess.run(
                [
                    "powershell",
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-Command",
                    ps_script,
                ],
                check=False,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
    except (OSError, subprocess.SubprocessError):
        # The main script remains usable if menu integration is unavailable.
        pass


install_application_menu_entry()


def ask_yes_no(prompt, default=None):
    while True:
        suffix = " [y/n]: " if default is None else (" [Y/n]: " if default else " [y/N]: ")
        answer = input(prompt + suffix).strip().lower()
        if not answer and default is not None:
            return default
        if answer in ("y", "yes"):
            return True
        if answer in ("n", "no"):
            return False
        print("Enter y or n.")


def confirm(prompt):
    return ask_yes_no(prompt + " This action may be difficult to undo.", False)


def read_int(prompt, minimum, maximum):
    while True:
        raw = input(prompt).strip()
        try:
            value = int(raw)
        except ValueError:
            print("Enter a whole number.")
            continue
        if minimum <= value <= maximum:
            return value
        print(f"Enter a number from {minimum} to {maximum}.")


def read_choice(prompt, choices):
    choices = {str(choice).lower() for choice in choices}
    while True:
        answer = input(prompt).strip().lower()
        if answer in choices:
            return answer
        print("Choose one of: " + ", ".join(sorted(choices)))


def show_disclaimer():
    print("\n" + DISCLAIMER)


def show_tutorial():
    print("\nPixelpass tutorial")
    print("1. Show tutorial")
    print("2. Show Disclaimer")
    print("3. Generating a password")
    print("4. Generating a passphrace")
    print("5. Generating a username")
    print("6. Generating new password")
    print("7. Viewing saved passwords")
    print("8. Editing saved passwords")
    print("9. Deleting saved passwords")
    print("10. Changing some settings by your comfort")
    print("0. Leave the script")
    print("\nDisclaimer:")
    print(DISCLAIMER)


def generate_password():
    print("\nGenerate password")
    length = read_int("Password length (1-99999): ", 1, MAX_INPUT_LENGTH)

    categories = []
    if ask_yes_no("Allow lowercase letters (a-z)?", True):
        categories.append(("lowercase", "abcdefghijklmnopqrstuvwxyz"))
    if ask_yes_no("Allow uppercase letters (A-Z)?", True):
        categories.append(("uppercase", "ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    if ask_yes_no("Allow digits (0-9)?", True):
        categories.append(("digits", "0123456789"))
    if ask_yes_no("Allow symbols?", True):
        categories.append(("symbols", "!@#$%^&*()-_=+[]{};:,.<>?/\\|~`'\""))

    extra = input("Extra characters to allow (Enter for none): ")
    if extra:
        categories.append(("custom characters", extra))

    if not categories:
        print("Select at least one character category.")
        return

    excluded_chars = input("Characters to exclude (Enter for none): ")
    excluded_terms = input(
        "Words or character sequences to exclude, comma-separated (Enter for none): "
    )
    forbidden = [
        term.strip().casefold()
        for term in excluded_terms.split(",")
        if term.strip()
    ]

    cleaned_categories = []
    for name, characters in categories:
        cleaned = "".join(
            dict.fromkeys(char for char in characters if char not in excluded_chars)
        )
        if cleaned:
            cleaned_categories.append((name, cleaned))
        else:
            print(f"The {name} category has no allowed characters left.")

    if not cleaned_categories:
        print("No allowed characters remain.")
        return

    if length < len(cleaned_categories):
        print(
            f"Length must be at least {len(cleaned_categories)} to include "
            "at least one character from every selected category."
        )
        return

    pool = "".join(dict.fromkeys("".join(chars for _, chars in cleaned_categories)))
    rng = secrets.SystemRandom()

    for _ in range(10000):
        result = [secrets.choice(chars) for _, chars in cleaned_categories]
        result.extend(secrets.choice(pool) for _ in range(length - len(result)))
        rng.shuffle(result)
        password = "".join(result)
        if not any(term in password.casefold() for term in forbidden):
            print("\nGenerated password:")
            print(password)
            return

    print("Could not satisfy those exclusions. Try fewer excluded words or sequences.")


WORDS = """
amber anchor apple arch artist autumn badger bamboo beacon berry birch bird
blossom blue breeze brook cabin cactus candle canyon cedar cherry circle cloud
coast cocoa comet coral cosmos cotton cove crystal daisy dawn delta diamond
drift dune eagle earth echo elm ember fabric falcon fern field firefly fjord
flame forest fox frost garden garlic gazelle glacier globe gold granite grape
grass harbor hazel heart honey horizon island ivory jasmine jigsaw jungle juniper
kestrel kitchen lagoon lantern lemon lilac lily lizard lotus maple marble meadow
melody meteor mint mirror mist moon morning mountain music navy nectar nickel
night noodle north ocean olive orbit orchid otter oyster palace paper pebble
pepper petal piano pickle pillow planet plum pocket poetry pond poppy prairie
puzzle quartz rabbit rainbow raven recipe reef river robin rocket rose ruby
saddle saffron sailor scarlet season shadow shell silver sky smoke snow solar
sparrow spice spring spruce star stone storm summer sunset swift temple thimble
thunder tiger timber toast tomato topaz torch tulip tunnel turtle valley velvet
violet walnut water wheat willow window winter wren yellow zephyr
acorn agate almond alpine anvil apricot arctic arrow atlas avenue basil basket
bayou beach beaver bedroom beetle bell bison blanket blueberry boardwalk bonfire
boulder bouquet brandy bridge bronze bubble buffalo butter button cadenza camel
camellia canary cardinal caribou carpenter carrot cascade castle cat cedar
celery chamber charcoal charm cheetah cinnamon citadel clover cobalt coconut
compass conifer copper coriander courtyard crane cranberry creek crescent cricket
crown custard cyclone cypress dagger daffodil daisywood dandelion daybreak denim
desert dewberry dragon drum dryad duchess dumpling earl east eclipse eggplant
eider elk emerald emperor enamel engine envelope espresso estuary evergreen falcon
feather festival finch fire firework flannel flour flower flute foothill fountain
foxglove frangipani freckle freesia freezer fresco frolic galena galaxy galleon
garnet gazelle gazebo geranium ginger ginkgo glen glitter goldfinch goose granite
grapefruit greenhouse grover grouse guitar gyroscope hail hammock handbell harvest
hatchery heather hedgehog helix hemlock heron hibiscus hilltop hippo hollyhock
homestead horseshoe hotcake hummingbird hyacinth iceberg inkling ironwood jackal
jade jaguar javelin jay jasmine jaybird jelly junebug kaleidoscope kayak keystone
kiwi koala kumquat labrador lacework ladybug lake lavender leaf leaping lemur
leopard lighthouse linen lion locket logbook lullaby magnolia magnolia mapleleaf
marigold marimba marmot marshmallow mastiff matinee mayflower memento meridian
midnight milkweed millstone minnow mirage mongoose moonglow morse moss mulberry
mustard narwhal nebula newt nightingale nimble nomad nutmeg oatmeal obsidian
octagon octopus odyssey opal orange orchard oregano origami osprey oxbow oyster
paddock paprika parakeet parsley parsnip pasture peacock peanut pelican pendulum
penguin persimmon phoenix pinecone ping pong pinewood pinwheel placid popcorn
porcupine prairie dog primrose prism promenade pumpkin quail quasar quill quince
radar radish ragtime raisin rattlesnake redwood rhubarb ribbon ricochet ringlet
rosemary rowboat sandalwood sapphire saucer seaglass seahorse sequoia sesame
shamrock shanty sherbet shingle shipwreck shoreline shuttle skylark slate
snowflake soapstone sorrel south spearmint spindle starling starlight stethoscope
strawberry sunrise surfboard swan sweetpea tadpole tapestry teacup teapot thicket
thrush tidepool tinsel toaster trillium truffle turnip twine umbrella unicorn
urchin valerian vanilla vanguard veranda vineyard vulture wagon wheelbarrow
whirlpool whistle wildflower windmill windswept winery woodpecker woodruff yarrow
yonder zinnia zircon
""".split()


def generate_passphrase():
    print("\nGenerate passphrase")
    count = read_int("Number of words (2-99999): ", 2, MAX_INPUT_LENGTH)
    separator = input("Separator between words (Enter for a space): ")
    if separator == "":
        separator = " "

    words = [secrets.choice(WORDS) for _ in range(count)]
    if ask_yes_no("Capitalize the first letter of each word?", False):
        words = [word.capitalize() for word in words]

    print("\nGenerated passphrase:")
    print(separator.join(words))


def generate_username():
    print("\nGenerate username")
    base = input("Optional name or word to include (Enter for random): ").strip()
    separator = input("Separator to use (Enter for underscore): ")
    if separator == "":
        separator = "_"

    style = read_choice(
        "Choose style: 1=random word + digits, 2=two random words, "
        "3=your word + random digits: ",
        ("1", "2", "3"),
    )

    if style == "1":
        word = secrets.choice(WORDS)
        digits = "".join(secrets.choice("0123456789") for _ in range(4))
        username = separator.join((word, digits))
    elif style == "2":
        username = separator.join((secrets.choice(WORDS), secrets.choice(WORDS)))
    else:
        if not base:
            base = input("Enter a name or word: ").strip()
        if not base:
            print("A name or word is required for this style.")
            return
        digits = "".join(secrets.choice("0123456789") for _ in range(4))
        username = separator.join((base, digits))

    print("\nGenerated username:")
    print(username)


def load_entries():
    ensure_vault()
    workbook = open_vault_workbook()
    try:
        sheet = workbook.active
        if sheet.max_row == 1 and all(
            sheet.cell(1, index + 1).value is None for index in range(4)
        ):
            for column, header in enumerate(HEADERS, start=1):
                sheet.cell(1, column, header)
            save_vault_workbook(workbook)

        actual_headers = [sheet.cell(1, index + 1).value for index in range(4)]
        if actual_headers != HEADERS:
            raise ValueError("The vault spreadsheet has unexpected columns.")

        entries = []
        for row in range(2, sheet.max_row + 1):
            values = [sheet.cell(row, col).value for col in range(1, 5)]
            if any(value is not None for value in values):
                entries.append(
                    (row, ["" if value is None else str(value) for value in values])
                )
        return entries
    finally:
        workbook.close()


def write_cell(cell, value):
    cell.value = str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def save_entry(row, values):
    workbook = open_vault_workbook()
    try:
        sheet = workbook.active
        for column, value in enumerate(values, start=1):
            write_cell(sheet.cell(row, column), value)
        save_vault_workbook(workbook)
    finally:
        workbook.close()


def append_entry(values):
    workbook = open_vault_workbook()
    try:
        sheet = workbook.active
        row = max(2, sheet.max_row + 1)
        for column, value in enumerate(values, start=1):
            write_cell(sheet.cell(row, column), value)
        save_vault_workbook(workbook)
    finally:
        workbook.close()


def delete_entry(row):
    workbook = open_vault_workbook()
    try:
        workbook.active.delete_rows(row, 1)
        save_vault_workbook(workbook)
    finally:
        workbook.close()


def create_new_entry():
    print("\nCreate new password")
    site = input("Site or custom name: ").strip()
    if not site:
        print("A site or custom name is required.")
        return

    try:
        entries = load_entries()
    except Exception as error:
        print(f"Could not read the vault: {error}")
        return

    if any(values[0].casefold() == site.casefold() for _, values in entries):
        print("That site or custom name already exists. Edit its saved entry instead.")
        return

    username = ""
    password = ""
    if ask_yes_no("Do you want to create a username?"):
        if ask_yes_no("Generate a username now?", True):
            generate_username()
            username = input("Username to save (paste or type it here): ")
        else:
            username = input("Username: ")

    email = input("Email (Enter to leave blank): ").strip()

    if ask_yes_no("Do you want to create a password?"):
        if ask_yes_no("Generate a password now?", True):
            generate_password()
            password = getpass.getpass("Password to save (paste or type it here): ")
        else:
            password = getpass.getpass("Password: ")

    if not username and not email and not password:
        print("Nothing to save.")
        return

    if not confirm("Save this entry to the vault?"):
        print("Entry not saved.")
        return

    try:
        append_entry([site, username, email, password])
        print("Entry saved.")
    except Exception as error:
        print(f"Could not save the entry: {error}")


def choose_entry():
    try:
        entries = load_entries()
    except Exception as error:
        print(f"Could not read the vault: {error}")
        return None

    if not entries:
        print("No saved entries.")
        return None

    for number, (_, values) in enumerate(entries, start=1):
        print(f"{number}. {values[0]}")

    selection = read_int("Choose an entry number (0 to cancel): ", 0, len(entries))
    if selection == 0:
        return None
    return entries[selection - 1]


def view_entries():
    try:
        entries = load_entries()
    except Exception as error:
        print(f"Could not read the vault: {error}")
        return

    if not entries:
        print("No saved entries.")
        return

    for number, (_, values) in enumerate(entries, start=1):
        print(f"{number}. {values[0]}")

    selection = read_int("Choose an entry number (0 to cancel): ", 0, len(entries))
    if selection == 0:
        return
    if not confirm("Reveal the selected entry, including its password?"):
        print("Password not revealed.")
        return

    _, values = entries[selection - 1]
    print(f"\nSite or custom name: {values[0]}")
    print(f"Username: {values[1]}")
    print(f"Email: {values[2]}")
    print(f"Password: {values[3]}")


def edit_entry():
    selected = choose_entry()
    if selected is None:
        return

    row, original = selected
    values = original.copy()
    print("Press Enter to keep a value unchanged.")
    for index, label in enumerate(HEADERS):
        replacement = input(f"{label} [{values[index]}]: ")
        if replacement != "":
            values[index] = replacement

    if values == original:
        print("No changes made.")
        return

    try:
        entries = load_entries()
        for other_row, other_values in entries:
            if other_row != row and other_values[0].casefold() == values[0].casefold():
                print("Another entry already uses that site or custom name.")
                return
    except Exception as error:
        print(f"Could not validate the changes: {error}")
        return

    if not confirm("Save changes to this entry?"):
        print("Changes not saved.")
        return

    try:
        save_entry(row, values)
        print("Entry updated.")
    except Exception as error:
        print(f"Could not update the entry: {error}")


def remove_entry():
    selected = choose_entry()
    if selected is None:
        return
    row, values = selected
    if not confirm(f"Delete the entry for {values[0]!r}?"):
        print("Entry not deleted.")
        return
    try:
        delete_entry(row)
        print("Entry deleted.")
    except Exception as error:
        print(f"Could not delete the entry: {error}")


def read_new_master_password():
    while True:
        password = getpass.getpass("New master password: ")
        if not password:
            print("The master password cannot be empty.")
            continue
        confirmation = getpass.getpass("Re-enter the new master password: ")
        if password != confirmation:
            print("The passwords do not match.")
            continue
        return password


def set_master_password():
    global MASTER_PASSWORD, ENCRYPTED_MODE

    if ENCRYPTED_MODE:
        print("A master password is already set. Choose Change or Remove instead.")
        return

    print(
        "\nWarning: if you forget the master password, the encrypted data "
        "will be inaccessible and cannot be recovered."
    )
    if not ask_yes_no("Set a master password and encrypt the vault?", False):
        return

    try:
        ensure_vault()
        workbook = open_vault_workbook()
        try:
            new_password = read_new_master_password()
            encrypted = encrypt_bytes(workbook_to_bytes(workbook), new_password)
            atomic_write(ENCRYPTED_VAULT_PATH, encrypted)
        finally:
            workbook.close()

        MASTER_PASSWORD = new_password
        ENCRYPTED_MODE = True
        try:
            VAULT_PATH.unlink(missing_ok=True)
        except OSError:
            print("Warning: the old unencrypted spreadsheet could not be removed.")
        SETTINGS["vault_file"] = "vault.xlsx"
        save_settings(SETTINGS)
        print("Master password set. The vault is now encrypted.")
    except Exception as error:
        print(f"Could not set the master password: {error}")


def change_master_password():
    global MASTER_PASSWORD

    if not ENCRYPTED_MODE or MASTER_PASSWORD is None:
        print("Set a master password first.")
        return

    if not ask_yes_no("Change the master password?", False):
        return

    try:
        workbook = open_vault_workbook()
        try:
            new_password = read_new_master_password()
            encrypted = encrypt_bytes(workbook_to_bytes(workbook), new_password)
            atomic_write(ENCRYPTED_VAULT_PATH, encrypted)
        finally:
            workbook.close()

        MASTER_PASSWORD = new_password
        print("Master password changed.")
    except Exception as error:
        print(f"Could not change the master password: {error}")


def remove_master_password():
    global MASTER_PASSWORD, ENCRYPTED_MODE

    if not ENCRYPTED_MODE or MASTER_PASSWORD is None:
        print("No master password is currently set.")
        return

    print("Removing the master password will leave the spreadsheet unencrypted.")
    if not ask_yes_no("Remove the master password?", False):
        return

    try:
        workbook = open_vault_workbook()
        try:
            workbook.save(VAULT_PATH)
        finally:
            workbook.close()

        ENCRYPTED_VAULT_PATH.unlink()
        MASTER_PASSWORD = None
        ENCRYPTED_MODE = False
        SETTINGS["vault_file"] = "vault.xlsx"
        save_settings(SETTINGS)
        print("Master password removed. The spreadsheet is now unencrypted.")
    except Exception as error:
        print(f"Could not remove the master password: {error}")


def master_password_menu():
    while True:
        print("\nMaster password")
        print("1. Set master password")
        print("2. Change master password")
        print("3. Remove master password")
        print("0. Back")
        choice = input("Choose an option: ").strip()

        if choice == "1":
            set_master_password()
        elif choice == "2":
            change_master_password()
        elif choice == "3":
            remove_master_password()
        elif choice == "0":
            return
        else:
            print("Enter a number from 0 to 3.")


def settings_menu():
    while True:
        clear_setting = SETTINGS.get("clear_terminal", "false") == "true"
        print("\nSettings")
        print(f"1. Clear terminal after each operation: {'On' if clear_setting else 'Off'}")
        print("2. Master password options")
        print("0. Back")
        choice = input("Choose an option: ").strip()

        if choice == "1":
            SETTINGS["clear_terminal"] = "false" if clear_setting else "true"
            save_settings(SETTINGS)
            print(
                "Terminal clearing is now "
                + ("on." if SETTINGS["clear_terminal"] == "true" else "off.")
            )
        elif choice == "2":
            master_password_menu()
        elif choice == "0":
            return
        else:
            print("Enter a number from 0 to 2.")


def clear_terminal():
    # ANSI clear sequence; does not invoke a shell or access the network.
    print("\033[2J\033[H", end="")


def main():
    global ENCRYPTED_MODE

    ensure_vault()
    if ENCRYPTED_VAULT_PATH.exists():
        if not unlock_encrypted_vault():
            print("Vault remains locked. Goodbye.")
            return

    print("welcome to pixelpass!")

    while True:
        print("\n--- pixelpass ---")
        print("1. Tutorial")
        print("2. Disclaimer")
        print("3. Generate password")
        print("4. Generate passphrase")
        print("5. Generate username")
        print("6. Create new password")
        print("7. View saved passwords")
        print("8. Edit saved password")
        print("9. Delete saved entry")
        print("10. Settings")
        print("0. Exit")
        choice = input("Enter the number of the action you want to do: ").strip()

        if choice == "1":
            show_tutorial()
        elif choice == "2":
            show_disclaimer()
        elif choice == "3":
            generate_password()
        elif choice == "4":
            generate_passphrase()
        elif choice == "5":
            generate_username()
        elif choice == "6":
            create_new_entry()
        elif choice == "7":
            view_entries()
        elif choice == "8":
            edit_entry()
        elif choice == "9":
            remove_entry()
        elif choice == "10":
            settings_menu()
        elif choice == "0":
            print("Goodbye.")
            break
        else:
            print("Enter a number from 0 to 10.")
            continue

        if SETTINGS.get("clear_terminal") == "true":
            input("\nPress Enter to continue...")
            clear_terminal()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nGoodbye.")
    except Exception as error:
        print(f"Pixelpass stopped because of an error: {error}")
